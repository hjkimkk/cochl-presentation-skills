import json, os, subprocess, sys, tempfile, unittest
from pathlib import Path

HERE = Path(__file__).parent
FX = HERE / "fixtures"

def run(*args):
    return subprocess.run([sys.executable, *args], capture_output=True, text=True)

def _probe():
    try:
        from playwright.sync_api import sync_playwright
        with sync_playwright() as p:
            b = p.chromium.launch()
            b.close()
        return True
    except Exception:
        return False

CHROMIUM_OK = _probe()

class FactsTest(unittest.TestCase):
    def facts(self, name):
        r = run(str(HERE / "facts.py"), str(FX / name))
        self.assertEqual(r.returncode, 0, r.stderr)
        return json.loads(r.stdout)

    @unittest.skipUnless(CHROMIUM_OK, "Chromium cannot launch here")
    def test_html_facts(self):
        f = self.facts("slide_a.html")
        self.assertEqual(f["kind"], "html")
        self.assertIn("IBM Plex Sans", f["fonts"])
        self.assertIn("#4b68ff", f["colors"])
        self.assertIn("#0b0b12", f["colors"])
        self.assertEqual(f["canvas"], [1440, 810])
        self.assertEqual(f["layouts"], ["hero"])
        self.assertLess(f["logos"][0]["distortion"], 0.05)

    @unittest.skipUnless(CHROMIUM_OK, "Chromium cannot launch here")
    def test_diff_flags_color_and_logo(self):
        a, b = HERE / "a.json", HERE / "b.json"
        a.write_text(json.dumps(self.facts("slide_a.html")))
        b.write_text(json.dumps(self.facts("slide_b.html")))
        try:
            r = run(str(HERE / "diff.py"), str(a), str(b))
            self.assertEqual(r.returncode, 1)
            self.assertIn("#832bfb", r.stdout)
            self.assertIn("logo", r.stdout.lower())
            same = run(str(HERE / "diff.py"), str(a), str(a))
            self.assertEqual(same.returncode, 0, same.stdout)
        finally:
            a.unlink(); b.unlink()

class DiffLogicTest(unittest.TestCase):
    def test_diff_logic_without_browser(self):
        base = {"kind": "html", "fonts": {"IBM Plex Sans": 2}, "colors": {"#4b68ff": 1, "#0b0b12": 1},
                "canvas": [1440, 810], "layouts": ["hero"],
                "logos": [{"src": "x", "distortion": 0.0}]}
        changed = json.loads(json.dumps(base))
        changed["colors"]["#832bfb"] = 1
        changed["logos"][0]["distortion"] = 1.0
        with tempfile.TemporaryDirectory() as d:
            a, b = Path(d) / "a.json", Path(d) / "b.json"
            a.write_text(json.dumps(base)); b.write_text(json.dumps(changed))
            r = run(str(HERE / "diff.py"), str(a), str(b))
            self.assertEqual(r.returncode, 1, r.stdout + r.stderr)
            self.assertIn("#832bfb", r.stdout)
            self.assertIn("logo", r.stdout.lower())
            same = run(str(HERE / "diff.py"), str(a), str(a))
            self.assertEqual(same.returncode, 0, same.stdout + same.stderr)
            self.assertIn("no brand-fact differences", same.stdout)

class AllowGateTest(unittest.TestCase):
    THEME = ":root { --accent: #4b68ff; --text: #fff; --font-sans: 'IBM Plex Sans', Arial, sans-serif; }\n"

    def gate(self, before, after, allow=True):
        def facts(colors, fonts):
            return {"kind": "html", "fonts": {f: 1 for f in fonts}, "colors": {c: 1 for c in colors},
                    "canvas": [1440, 810], "layouts": ["hero"], "logos": []}
        b = facts(before.get("colors", ["#4b68ff"]), before.get("fonts", ["IBM Plex Sans"]))
        a = facts(after.get("colors", before.get("colors", ["#4b68ff"])),
                  after.get("fonts", before.get("fonts", ["IBM Plex Sans"])))
        with tempfile.TemporaryDirectory() as d:
            pb, pa, pt = Path(d) / "b.json", Path(d) / "a.json", Path(d) / "theme.md"
            pb.write_text(json.dumps(b)); pa.write_text(json.dumps(a)); pt.write_text(self.THEME)
            args = [str(HERE / "diff.py"), str(pb), str(pa)] + (["--allow", str(pt)] if allow else [])
            return run(*args)

    def test_theme_color_swap_passes(self):
        r = self.gate({"colors": ["#4b68ff", "#ececec"]}, {"colors": ["#4b68ff", "#ffffff"]})
        self.assertEqual(r.returncode, 0, r.stdout + r.stderr)

    def test_off_theme_added_color_fails(self):
        r = self.gate({"colors": ["#4b68ff"]}, {"colors": ["#4b68ff", "#123456"]})
        self.assertEqual(r.returncode, 1, r.stdout)
        self.assertIn("#123456", r.stdout)
        self.assertIn("off-theme", r.stdout)

    def test_missing_theme_color_fails(self):
        r = self.gate({"colors": ["#4b68ff", "#ececec"]}, {"colors": ["#ececec"]})
        self.assertEqual(r.returncode, 1, r.stdout)
        self.assertIn("#4b68ff", r.stdout)

    def test_fonts(self):
        bad = self.gate({}, {"fonts": ["IBM Plex Sans", "Comic Sans MS"]})
        self.assertEqual(bad.returncode, 1, bad.stdout)
        self.assertIn("Comic Sans MS", bad.stdout)
        ok = self.gate({}, {"fonts": ["IBM Plex Sans", "Arial"]})
        self.assertEqual(ok.returncode, 0, ok.stdout)

    def test_without_allow_still_strict(self):
        r = self.gate({"colors": ["#4b68ff", "#ececec"]}, {"colors": ["#4b68ff", "#ffffff"]}, allow=False)
        self.assertEqual(r.returncode, 1, r.stdout)

class PptxFactsTest(unittest.TestCase):
    def test_pptx_facts(self):
        from pptx import Presentation
        from pptx.dml.color import RGBColor
        from pptx.util import Inches
        prs = Presentation()
        prs.slide_width, prs.slide_height = Inches(10), Inches(5.625)
        slide = prs.slides.add_slide(prs.slide_layouts[6])
        tb = slide.shapes.add_textbox(Inches(1), Inches(1), Inches(4), Inches(1))
        run_ = tb.text_frame.paragraphs[0].add_run()
        run_.text = "Hello"
        run_.font.name = "Century Gothic"
        run_.font.color.rgb = RGBColor(0x83, 0x2B, 0xFB)
        with tempfile.TemporaryDirectory() as d:
            path = Path(d) / "deck.pptx"
            prs.save(str(path))
            r = run(str(HERE / "facts.py"), str(path))
        self.assertEqual(r.returncode, 0, r.stderr)
        f = json.loads(r.stdout)
        self.assertEqual(f["canvas"], [10.0, 5.625])
        self.assertIn("Century Gothic", f["fonts"])
        self.assertIn("#832bfb", f["colors"])

if __name__ == "__main__":
    unittest.main()
