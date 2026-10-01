#!/usr/bin/env python3
"""Extract brand facts (fonts, colors, canvas, logo distortion, layouts) from an HTML or PPTX output."""
import json, sys
from collections import Counter
from pathlib import Path

def html_facts(path: Path, shots: Path | None):
    from playwright.sync_api import sync_playwright
    js = Path(__file__).with_name("facts.js").read_text()
    with sync_playwright() as p:
        b = p.chromium.launch()
        page = b.new_page(viewport={"width": 1440, "height": 810})
        page.goto(path.resolve().as_uri(), wait_until="networkidle")
        page.evaluate("document.fonts.ready")
        facts = page.evaluate(js)
        if shots:
            shots.mkdir(parents=True, exist_ok=True)
            els = page.query_selector_all("section.slide, .slide, .pg, .page") or [page]
            for i, el in enumerate(els, 1):
                el.screenshot(path=str(shots / f"{i:02d}.png"))
        b.close()
    return {"kind": "html", **facts}

def pptx_facts(path: Path):
    from pptx import Presentation
    prs = Presentation(str(path))
    fonts, colors = Counter(), Counter()
    for slide in prs.slides:
        for shape in slide.shapes:
            if not shape.has_text_frame:
                continue
            for para in shape.text_frame.paragraphs:
                for run in para.runs:
                    if run.font.name:
                        fonts[run.font.name] += 1
                    try:
                        if run.font.color and run.font.color.type is not None and run.font.color.rgb:
                            colors["#" + str(run.font.color.rgb).lower()] += 1
                    except AttributeError:
                        pass  # theme colors have no .rgb
    emu_per_in = 914400
    return {
        "kind": "pptx", "fonts": dict(fonts), "colors": dict(colors), "logos": [],
        "canvas": [round(prs.slide_width / emu_per_in, 3), round(prs.slide_height / emu_per_in, 3)],
        "layouts": [s.slide_layout.name for s in prs.slides],
    }

def main(argv):
    path = Path(argv[0])
    shots = Path(argv[argv.index("--shots") + 1]) if "--shots" in argv else None
    facts = pptx_facts(path) if path.suffix == ".pptx" else html_facts(path, shots)
    print(json.dumps(facts, indent=2, sort_keys=True))

if __name__ == "__main__":
    main(sys.argv[1:])
